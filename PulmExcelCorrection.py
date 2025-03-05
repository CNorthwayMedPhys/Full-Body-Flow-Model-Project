# -*- coding: utf-8 -*-
"""
Created on Tue Mar  4 14:29:53 2025

@author: Cassidy.Northway
"""

import tkinter as tk
import tkinter.filedialog as fd
import numpy as np
import easygui
import openpyxl

#%%
#load in an excel file
tk.Tk().withdraw()
file = fd.askopenfilename()
wb = openpyxl.load_workbook(file,data_only=True)
ws = wb.worksheets[0]

#%% Renormalize the time values
time_col = ws['A']
otime = []
for cell in time_col:
    otime.append(cell.value)
    
#Determine the size of the time array and create a [0,0.955] s array of the same size
maxtime = 0.955
time_size = len(otime)-1
ntime = [ws['A1'].value]
otime = otime[1:]

for value in otime:
    time = ((value-min(otime))*0.955)/(max(otime)-min(otime))
    ntime=np.append(ntime,time)

i=1
for value in ntime:
    ws.cell(i,1).value=value
    i += 1

#%%

#Pull Q values from excel
q_col = ws['B']
q_array = []
for cell in q_col:
        q_array.append(cell.value)
q_array = q_array[1:] # Drop B1

q_size = len(q_array)

#Ask for R Distal, R Prox and  Vessel Length cm
L = float(easygui.enterbox('Enter the vessel length in cm')) #cm
Rdist = float(easygui.enterbox('Enter the distal diam in cm'))/2 #cm
Rprox = float(easygui.enterbox('Enter the proximal diam in cm'))/2 #cm

#Calc R increments along the length of the vessel uses exp
X = np.linspace(0.0, L , 20) #cm


i = 2
#Save the new length array to the excel sheet
for x in X:
    ws.cell(1,i).value = x/100 # m 
    i += 1 

Rarray = Rdist * np.power((Rprox/Rdist), X/L)#cm

#Calc Area increments (cm^2)
Aarray = np.pi* Rarray**2 #cm^2
#Calculate velocity for each physical step and write to Excel
j=2
for A in Aarray:
    Varray = q_array/A #cm/s
    Varray = Varray/100 #m/s
    i=2
    for value in Varray:
        ws.cell(i,j).value = value
        i+=1
    j += 1

#Write, save, and close the excel sheet    
wb.save(file)
wb.close()